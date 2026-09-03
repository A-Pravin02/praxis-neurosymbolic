"""Dataset-agnostic parsing: bytes/records in, normalised documents out.

A Parser turns a source into one or more `ParsedDocument`s. It must not assume
anything about the *meaning* of the content — that is the adapter's job at
ingestion time. Adding a new format means adding a Parser, not changing the
pipeline.
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol

from app.core.errors import UnsupportedMediaError, ValidationError


@dataclass
class ParsedSection:
    """A heading-delimited region of a document."""
    heading: str
    text: str
    page: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    title: str
    sections: list[ParsedSection]
    media_type: str = "text/plain"
    metadata: dict[str, Any] = field(default_factory=dict)
    # Structured rows, when the source is tabular/record oriented. These feed the
    # knowledge extractor directly rather than going through text chunking.
    records: list[dict[str, Any]] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(
            (f"{s.heading}\n{s.text}" if s.heading else s.text) for s in self.sections
        ).strip()


class Parser(Protocol):
    name: str
    extensions: tuple[str, ...]
    media_types: tuple[str, ...]

    def parse(self, data: bytes, *, filename: str = "",
              media_type: str = "") -> ParsedDocument: ...


# ---------------------------------------------------------------------------

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*\S)\s*$", re.M)
_UNDERLINE_HEADING = re.compile(r"^(?P<title>[^\n]{3,80})\n[=-]{3,}\s*$", re.M)
_NUMBERED_HEADING = re.compile(
    r"^(?P<num>\d+(?:\.\d+)*\.?)\s+(?P<title>[A-Z][^\n]{2,80})\s*$", re.M)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def split_sections(text: str) -> list[ParsedSection]:
    """Split on markdown, underlined or numbered headings; fall back to one section."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    markers: list[tuple[int, int, str]] = []
    for m in _MD_HEADING.finditer(text):
        markers.append((m.start(), m.end(), m.group(2).strip()))
    for m in _UNDERLINE_HEADING.finditer(text):
        markers.append((m.start(), m.end(), m.group("title").strip()))
    for m in _NUMBERED_HEADING.finditer(text):
        markers.append((m.start(), m.end(), f"{m.group('num')} {m.group('title')}".strip()))
    markers.sort()
    if not markers:
        body = text.strip()
        return [ParsedSection(heading="", text=body)] if body else []

    sections: list[ParsedSection] = []
    preamble = text[: markers[0][0]].strip()
    if preamble:
        sections.append(ParsedSection(heading="", text=preamble))
    for i, (_, end, heading) in enumerate(markers):
        stop = markers[i + 1][0] if i + 1 < len(markers) else len(text)
        body = text[end:stop].strip()
        if body or heading:
            sections.append(ParsedSection(heading=heading, text=body))
    return sections


class TextParser:
    name = "text"
    extensions = (".txt", ".md", ".markdown", ".rst", ".log")
    media_types = ("text/plain", "text/markdown", "text/x-markdown")

    def parse(self, data: bytes, *, filename: str = "",
              media_type: str = "") -> ParsedDocument:
        text = _decode(data)
        sections = split_sections(text)
        title = _title_from(sections, filename)
        return ParsedDocument(title=title, sections=sections,
                              media_type=media_type or "text/plain")


class PdfParser:
    name = "pdf"
    extensions = (".pdf",)
    media_types = ("application/pdf",)

    def parse(self, data: bytes, *, filename: str = "",
              media_type: str = "") -> ParsedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise UnsupportedMediaError("PDF support is not installed.") from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [(i + 1, (page.extract_text() or "").strip())
                     for i, page in enumerate(reader.pages)]
        except Exception as exc:
            raise ValidationError("The PDF could not be read.") from exc
        sections: list[ParsedSection] = []
        for number, page_text in pages:
            if not page_text:
                continue
            for section in split_sections(page_text) or [ParsedSection("", page_text)]:
                section.page = number
                sections.append(section)
        meta_title = ""
        try:
            meta_title = (reader.metadata.title or "").strip() if reader.metadata else ""
        except Exception:
            meta_title = ""
        title = meta_title or Path(filename).stem or "Untitled PDF"
        return ParsedDocument(title=title, sections=sections, media_type="application/pdf",
                              metadata={"page_count": len(pages)})


class DocxParser:
    name = "docx"
    extensions = (".docx",)
    media_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    def parse(self, data: bytes, *, filename: str = "",
              media_type: str = "") -> ParsedDocument:
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise UnsupportedMediaError("DOCX support is not installed.") from exc
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise ValidationError("The DOCX file could not be read.") from exc

        sections: list[ParsedSection] = []
        heading = ""
        buffer: list[str] = []

        def flush() -> None:
            body = "\n".join(buffer).strip()
            if body or heading:
                sections.append(ParsedSection(heading=heading, text=body))

        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            style = (paragraph.style.name or "") if paragraph.style else ""
            if style.lower().startswith("heading") and text:
                flush()
                heading, buffer = text, []
            elif text:
                buffer.append(text)
        flush()

        for table in document.tables:
            rows = [" | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows]
            if rows:
                sections.append(ParsedSection(heading="Table", text="\n".join(rows)))

        title = Path(filename).stem or "Untitled document"
        return ParsedDocument(title=title, sections=sections,
                              media_type=self.media_types[0])


class JsonParser:
    name = "json"
    extensions = (".json", ".jsonl", ".ndjson")
    media_types = ("application/json", "application/x-ndjson")

    def parse(self, data: bytes, *, filename: str = "",
              media_type: str = "") -> ParsedDocument:
        raw = _decode(data).strip()
        records: list[dict[str, Any]] = []
        if not raw:
            raise ValidationError("The JSON document is empty.")
        if filename.endswith((".jsonl", ".ndjson")) or raw.lstrip()[0] != "[" and "\n" in raw:
            for line in raw.splitlines():
                line = line.strip()
                if not line:
                    continue
                try:
                    value = json.loads(line)
                except json.JSONDecodeError:
                    records = []
                    break
                records.append(value if isinstance(value, dict) else {"value": value})
        if not records:
            try:
                value = json.loads(raw)
            except json.JSONDecodeError as exc:
                raise ValidationError("The JSON document could not be parsed.") from exc
            if isinstance(value, list):
                records = [v if isinstance(v, dict) else {"value": v} for v in value]
            elif isinstance(value, dict):
                records = [value]
            else:
                records = [{"value": value}]

        sections = [
            ParsedSection(heading=f"Record {i + 1}", text=_render_record(record))
            for i, record in enumerate(records[:2000])
        ]
        return ParsedDocument(title=Path(filename).stem or "JSON dataset",
                              sections=sections, media_type="application/json",
                              records=records)


class CsvParser:
    name = "csv"
    extensions = (".csv", ".tsv")
    media_types = ("text/csv", "text/tab-separated-values")

    def parse(self, data: bytes, *, filename: str = "",
              media_type: str = "") -> ParsedDocument:
        raw = _decode(data)
        delimiter = "\t" if filename.endswith(".tsv") else _sniff_delimiter(raw)
        reader = csv.DictReader(io.StringIO(raw), delimiter=delimiter)
        records = [
            {(k or "").strip(): (v or "").strip() for k, v in row.items() if k}
            for row in reader
        ]
        if not records:
            raise ValidationError("The CSV file has no data rows.")
        sections = [
            ParsedSection(heading=f"Row {i + 1}", text=_render_record(record))
            for i, record in enumerate(records[:5000])
        ]
        return ParsedDocument(title=Path(filename).stem or "CSV dataset",
                              sections=sections, media_type="text/csv",
                              records=records,
                              metadata={"columns": list(records[0].keys())})


def _sniff_delimiter(raw: str) -> str:
    sample = "\n".join(raw.splitlines()[:5])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _render_record(record: dict[str, Any]) -> str:
    parts = []
    for key, value in record.items():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        parts.append(f"{key}: {value}")
    return "\n".join(parts)


def _title_from(sections: list[ParsedSection], filename: str) -> str:
    for section in sections:
        if section.heading:
            return section.heading[:400]
    stem = Path(filename).stem
    if stem:
        return stem[:400]
    if sections:
        return sections[0].text.strip().splitlines()[0][:120] or "Untitled document"
    return "Untitled document"


PARSERS: list[Parser] = [
    TextParser(), PdfParser(), DocxParser(), JsonParser(), CsvParser(),
]

SUPPORTED_EXTENSIONS = tuple(sorted({e for p in PARSERS for e in p.extensions}))


def select_parser(filename: str = "", media_type: str = "") -> Parser:
    suffix = Path(filename).suffix.lower()
    for parser in PARSERS:
        if suffix and suffix in parser.extensions:
            return parser
    normalized = (media_type or "").split(";")[0].strip().lower()
    for parser in PARSERS:
        if normalized and normalized in parser.media_types:
            return parser
    if not suffix and not normalized:
        return PARSERS[0]
    raise UnsupportedMediaError(
        f"Unsupported file type. Supported extensions: {', '.join(SUPPORTED_EXTENSIONS)}."
    )


def parse_bytes(data: bytes, *, filename: str = "", media_type: str = "") -> ParsedDocument:
    return select_parser(filename, media_type).parse(
        data, filename=filename, media_type=media_type
    )
